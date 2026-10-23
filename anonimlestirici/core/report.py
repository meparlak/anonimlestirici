"""Eşleştirme tablosu (Excel) yazma ve okuma."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import __version__
from .types import Candidate, EntityType

if TYPE_CHECKING:
    from .engine import Session

WARNING = "Bu dosya gerçek kişisel verileri içerir. Yapay zekâya vermeyin, paylaşmayın. İşiniz bitince silin."
MAP_SHEET = "Eşleştirme"
HEADERS = ["Kod", "Tür", "Orijinal değer", "Diğer yazımlar", "Adet", "Dosya(lar)", "Konumlar", "Tespit yöntemi"]

_BOLD = Font(bold=True)
_WARN_FILL = PatternFill("solid", fgColor="FDECC8")
_HEAD_FILL = PatternFill("solid", fgColor="E4EAF4")


def _locations(c: Candidate, limit: int = 12) -> str:
    locs = []
    for o in c.occurrences:
        if o.location not in locs:
            locs.append(o.location)
    more = f" (+{len(locs) - limit})" if len(locs) > limit else ""
    return ", ".join(locs[:limit]) + more


def _sheet_rows(ws, rows: list[list], widths: list[int]) -> None:
    ws.append(HEADERS)
    for cell in ws[ws.max_row]:
        cell.font = _BOLD
        cell.fill = _HEAD_FILL
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=ws.max_row - len(rows) + 1, column=1)


def _row(c: Candidate, code: str) -> list:
    others = sorted(v for v in c.variants if v != c.value)
    return [
        code,
        c.etype.label,
        c.value,
        ", ".join(others),
        c.count,
        ", ".join(c.files),
        _locations(c),
        ", ".join(sorted(c.sources)),
    ]


def write_mapping(path: Path, accepted: list[Candidate], skipped: list[Candidate], session: "Session", files: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = MAP_SHEET
    ws.append([WARNING])
    ws["A1"].font = Font(bold=True, color="7A4B00")
    ws["A1"].fill = _WARN_FILL
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.append([])
    rows = [_row(c, c.code) for c in accepted if c.code]
    _sheet_rows(ws, rows, [14, 14, 32, 32, 8, 28, 36, 26])

    sk = wb.create_sheet("Atlananlar")
    sk.append(["Önizlemede onayı kaldırılan, dosyada olduğu gibi bırakılan değerler."])
    sk.append([])
    _sheet_rows(sk, [_row(c, "") for c in skipped], [14, 14, 32, 32, 8, 28, 36, 26])

    oz = wb.create_sheet("Özet")
    counts: dict[EntityType, int] = {}
    for c in accepted:
        counts[c.etype] = counts.get(c.etype, 0) + 1
    s = session.settings
    summary = [
        ["Uyarı", WARNING],
        ["İşlem tarihi", datetime.now().strftime("%d.%m.%Y %H:%M")],
        ["Uygulama", f"Anonimleştirici {__version__}"],
        ["İşlenen dosyalar", ", ".join(session.processed or files)],
        ["Kod biçimi", s.code_format],
        ["Açık türler", ", ".join(et.label for et in EntityType if s.is_enabled(et))],
        ["Özel liste", f"{len(s.custom_entries())} terim"],
        ["Asla anonimleştirme listesi", f"{len(s.exclude_entries())} terim"],
        [],
        ["Tür", "Farklı değer"],
    ] + [[et.label, n] for et, n in counts.items()]
    for r in summary:
        oz.append(r)
    oz.column_dimensions["A"].width = 28
    oz.column_dimensions["B"].width = 90
    for row in oz.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    oz["A1"].font = _BOLD
    wb.save(path)


def read_mapping(path: Path) -> dict[str, str]:
    """Eşleştirme tablosundan {kod: orijinal değer} sözlüğü döndürür."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if MAP_SHEET not in wb.sheetnames:
        raise ValueError(f"Bu dosyada \"{MAP_SHEET}\" sayfası yok. Anonimleştirici'nin ürettiği eşleştirme tablosunu seçin.")
    ws = wb[MAP_SHEET]
    out: dict[str, str] = {}
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not header_seen:
            header_seen = bool(row and row[0] == "Kod")
            continue
        if row and row[0] and row[2] is not None:
            out[str(row[0])] = str(row[2])
    wb.close()
    return out
