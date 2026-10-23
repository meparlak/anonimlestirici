"""Testler için küçük örnek dosyalar üretir. Bütün isimler ve numaralar uydurmadır."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

CORE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
 xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:creator>Kemal Kılıç</dc:creator>
<cp:lastModifiedBy>Kemal Kılıç</cp:lastModifiedBy></cp:coreProperties>"""


def make_docx(path: Path) -> Path:
    doc = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W}"><w:body>
<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Sayın Dr. Elif </w:t></w:r><w:r><w:t>De</w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t>mir,</w:t></w:r></w:p>
<w:p><w:r><w:t xml:space="preserve">Bilgi İşlem Daire Başkanlığı'nın yazısı ekte. İletişim: elif.demir@ornek.edu.tr, TC: 10000000078.</w:t></w:r></w:p>
<w:p><w:ins w:author="Murat Şahin"><w:r><w:t>Esentepe Mah. Üniversite Cad. No:12 Serdivan/Sakarya</w:t></w:r></w:ins></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Elif Demir</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>0532 123 45 67</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
</w:body></w:document>"""
    header = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>Hazırlayan: Kemal Kılıç</w:t></w:r></w:p></w:hdr>"""
    comments = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="{W}"><w:comment w:id="0" w:author="Murat Şahin" w:initials="MŞ"><w:p><w:r><w:t>Elif Hanım'a sorulsun.</w:t></w:r></w:p></w:comment></w:comments>"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", doc)
        z.writestr("word/header1.xml", header)
        z.writestr("word/comments.xml", comments)
        z.writestr("word/media/image1.png", b"\x89PNG fake")
        z.writestr("docProps/core.xml", CORE)
    return path


def make_pptx(path: Path) -> Path:
    slide = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree><p:sp><p:txBody>
<a:p><a:r><a:t>Proje sorumlusu: Ayşe </a:t></a:r><a:r><a:t>Kaya</a:t></a:r></a:p>
<a:p><a:r><a:t>Yılmaz Makine A.Ş. ile görüşüldü.</a:t></a:r></a:p>
</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"""
    notes = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:notes xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree><p:sp><p:txBody>
<a:p><a:r><a:t>Ayşe Kaya'yı 0532 123 45 67 numarasından ara.</a:t></a:r></a:p>
</p:txBody></p:sp></p:spTree></p:cSld></p:notes>"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/slides/slide1.xml", slide)
        z.writestr("ppt/notesSlides/notesSlide1.xml", notes)
        z.writestr("docProps/core.xml", CORE)
    return path


def make_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Şikayetler"
    ws.append(["Tarih", "Ad Soyad", "Telefon", "Birim", "Şikâyet"])
    ws.append(["03.10.2026", "Ahmet Yılmaz", 5321234567, "Satın Alma Birimi",
               "Ahmet Bey siparişin geciktiğini, Ayşe Kaya'ya ulaşamadığını söyledi."])
    ws.append(["05.10.2026", "AYŞE KAYA", "0264 295 10 10", "Kalite Müdürlüğü",
               "Yılmaz Makine A.Ş. teslimatında 3 hatalı parça."])
    ws.append(["06.10.2026", "Zehra Nur Yücel", "0505 000 00 00", "Kalite Müdürlüğü", "Ahmet Yılmaz tekrar aradı."])
    ws["E3"].comment = Comment("Mustafa Öztürk ile konuşuldu", "Kemal Kılıç")
    ws["F2"] = "=B2&\" - \"&D2"
    ozet = wb.create_sheet("Özet")
    ozet["A1"] = "Toplam"
    ozet["B1"] = "=COUNTA('Şikayetler'!B2:B4)"
    wb.properties.creator = "Kemal Kılıç"
    wb.save(path)
    return path


def make_csv(path: Path) -> Path:
    text = "Ad Soyad;TC;E-posta;Not\r\nAhmet Yılmaz;10000000078;ahmet@ornek.com;Ahmet Bey aradı\r\nŞükrü Öztürk;12345678901;sukru@ornek.com;\r\n"
    path.write_bytes(text.encode("cp1254"))
    return path


def make_txt(path: Path) -> Path:
    path.write_text(
        "Toplantı notu\nKatılanlar: Ahmet Yılmaz, Elif Demir (Kalite Müdürlüğü)\nAhmet Bey konuyu açtı. Tel: 0532 123 45 67\n",
        encoding="utf-8",
    )
    return path
