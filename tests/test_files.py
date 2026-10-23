import zipfile

from openpyxl import load_workbook

from anonimlestirici.core.engine import Session
from anonimlestirici.core.report import read_mapping
from anonimlestirici.core.reverse import reverse_file, reverse_text
from anonimlestirici.core.types import EntityType as T

from .fixtures import make_csv, make_docx, make_pptx, make_txt, make_xlsx


def run(settings, *paths, out_dir=None):
    s = Session(settings)
    result = s.analyze(list(paths))
    return s, result, s.save(result, out_dir)


def values(result):
    return {(c.etype, c.value) for c in result.candidates if c.accepted}


def zip_text(path, part):
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf-8")


# ------------------------------------------------------------------ Word


def test_docx_split_runs_and_formatting_kept(settings, tmp_path):
    src = make_docx(tmp_path / "yazi.docx")
    _, result, res = run(settings, src)
    assert (T.KISI, "Elif Demir") in values(result)
    doc = zip_text(res.outputs[0], "word/document.xml")
    assert "Elif" not in doc and "Demir" not in doc
    # İsim üç parçaya bölünmüştü; kalın ve italik biçimler hâlâ yerinde
    assert "<w:b/>" in doc and "<w:i/>" in doc
    assert "[KİŞİ_1]" in doc
    assert "elif.demir@ornek.edu.tr" not in doc and "10000000078" not in doc
    assert "Esentepe" not in doc and "0532 123 45 67" not in doc


def test_docx_header_comments_metadata(settings, tmp_path):
    src = make_docx(tmp_path / "yazi.docx")
    _, _, res = run(settings, src)
    out = res.outputs[0]
    assert "Kemal" not in zip_text(out, "word/header1.xml")
    comments = zip_text(out, "word/comments.xml")
    assert "Murat" not in comments and "Elif" not in comments and 'w:initials=""' in comments
    assert "Murat" not in zip_text(out, "word/document.xml")  # değişiklik izleme yazarı
    assert "Kemal" not in zip_text(out, "docProps/core.xml")
    # resim dokunulmadan kopyalandı
    assert zipfile.ZipFile(out).read("word/media/image1.png") == b"\x89PNG fake"


def test_original_file_untouched(settings, tmp_path):
    src = make_docx(tmp_path / "yazi.docx")
    before = src.read_bytes()
    run(settings, src)
    assert src.read_bytes() == before


# ------------------------------------------------------------------ Excel


def test_xlsx_cells_headers_and_numbers(settings, tmp_path):
    src = make_xlsx(tmp_path / "sikayetler.xlsx")
    _, result, res = run(settings, src)
    vals = values(result)
    assert (T.KISI, "Zehra Nur Yücel") in vals  # sütun başlığından
    wb = load_workbook(res.outputs[0])
    ws = wb["Şikayetler"]
    assert ws["B2"].value == "[KİŞİ_1]"
    assert ws["B3"].value == ws["E2"].value.split("'")[0].split(", ")[-1]  # AYŞE KAYA = Ayşe Kaya
    assert ws["C2"].value.startswith("[TELEFON_")  # sayı olarak saklanmış telefon
    assert "Ahmet" not in ws["E2"].value and "[KİŞİ_1] Bey" in ws["E2"].value
    assert ws["E4"].value == "[KİŞİ_1] tekrar aradı."
    assert ws["A2"].value == "03.10.2026"  # tarih türü kapalı
    assert ws["F2"].value == '=B2&" - "&D2'  # formül korunur
    assert "Mustafa" not in ws["E3"].comment.text
    assert wb.properties.creator != "Kemal Kılıç"


def test_xlsx_mapping_table(settings, tmp_path):
    src = make_xlsx(tmp_path / "sikayetler.xlsx")
    _, result, res = run(settings, src)
    assert res.mapping.name == "sikayetler_eslestirme.xlsx"
    mapping = read_mapping(res.mapping)
    assert mapping["[KİŞİ_1]"] == "Ahmet Yılmaz"
    wb = load_workbook(res.mapping)
    assert wb.sheetnames == ["Eşleştirme", "Atlananlar", "Özet"]
    assert "kişisel veri" in wb["Eşleştirme"]["A1"].value


def test_exclude_list_keeps_value(settings, tmp_path):
    settings.exclude_terms = ["Kalite Müdürlüğü"]
    src = make_xlsx(tmp_path / "sikayetler.xlsx")
    _, _, res = run(settings, src)
    ws = load_workbook(res.outputs[0])["Şikayetler"]
    assert ws["D3"].value == "Kalite Müdürlüğü"
    assert ws["D2"].value == "[BİRİM_1]"


def test_unchecked_candidate_is_left_alone(settings, tmp_path):
    src = make_xlsx(tmp_path / "sikayetler.xlsx")
    s = Session(settings)
    result = s.analyze([src])
    for c in result.candidates:
        if c.etype == T.BIRIM:
            c.accepted = False
    s.codebook.assign(result.candidates)
    res = s.save(result)
    ws = load_workbook(res.outputs[0])["Şikayetler"]
    assert ws["D2"].value == "Satın Alma Birimi"
    assert load_workbook(res.mapping)["Atlananlar"].max_row > 3


def test_sheet_name_with_person_is_renamed_and_formulas_follow(settings, tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Ahmet Yılmaz"
    wb.active["A1"] = "Ad Soyad"
    wb.active["B1"] = "Not"
    wb.active["A2"] = "Ahmet Yılmaz"
    other = wb.create_sheet("Toplam")
    other["A1"] = "='Ahmet Yılmaz'!A2"
    src = tmp_path / "kisi.xlsx"
    wb.save(src)
    _, _, res = run(settings, src)
    out = load_workbook(res.outputs[0])
    assert "KİŞİ_1" in out.sheetnames
    assert out["Toplam"]["A1"].value == "='KİŞİ_1'!A2"


# ------------------------------------------------------------------ PowerPoint, CSV, metin


def test_pptx(settings, tmp_path):
    src = make_pptx(tmp_path / "sunum.pptx")
    _, result, res = run(settings, src)
    slide = zip_text(res.outputs[0], "ppt/slides/slide1.xml")
    notes = zip_text(res.outputs[0], "ppt/notesSlides/notesSlide1.xml")
    assert "Kaya" not in slide and "Yılmaz Makine" not in slide
    assert "Ayşe" not in notes and "0532" not in notes


def test_csv_keeps_encoding_and_delimiter(settings, tmp_path):
    src = make_csv(tmp_path / "liste.csv")
    _, result, res = run(settings, src)
    raw = res.outputs[0].read_bytes()
    text = raw.decode("cp1254")
    assert text.startswith("Ad Soyad;TC;E-posta;Not\r\n")
    assert "Ahmet" not in text and "Şükrü" not in text
    assert "12345678901" not in text  # TC sütununda olduğu için doğrulama tutmasa da alınır
    assert "[KİŞİ_1] Bey aradı" in text


def test_txt_roundtrip_and_reverse(settings, tmp_path):
    src = make_txt(tmp_path / "not.txt")
    _, result, res = run(settings, src)
    out = res.outputs[0].read_text(encoding="utf-8")
    assert "Ahmet" not in out and "Elif" not in out and "0532" not in out
    back, stats = reverse_text(out, res.mapping)
    # "Ahmet Bey" aynı kişinin kısa yazımıydı; geri çevrilince tam adı alır
    assert back == src.read_text(encoding="utf-8").replace("Ahmet Bey", "Ahmet Yılmaz Bey")
    assert stats.replaced > 0 and not stats.unknown
    assert res.outputs[0].read_bytes()[:3] != b"\xef\xbb\xbf"  # kaynakta BOM yoksa çıktıda da yok


def test_reverse_tolerates_ai_edits(settings, tmp_path):
    src = make_txt(tmp_path / "not.txt")
    _, _, res = run(settings, src)
    text = "Kök neden: KİŞİ_1 ile [Kisi_2] arasında iletişim eksikliği. [KİŞİ_9] bilinmiyor."
    back, stats = reverse_text(text, res.mapping)
    assert "Ahmet Yılmaz" in back and "Elif Demir" in back
    assert "[KİŞİ_9]" in back and "[KİŞİ_9]" in stats.unknown


def test_reverse_docx_file(settings, tmp_path):
    src = make_docx(tmp_path / "yazi.docx")
    _, _, res = run(settings, src)
    back, _ = reverse_file(res.outputs[0], res.mapping)
    assert "Elif Demir" in zip_text(back, "word/document.xml")


def test_batch_shares_codes_across_files(settings, tmp_path):
    a = make_txt(tmp_path / "a.txt")
    b = make_xlsx(tmp_path / "b.xlsx")
    s, result, res = run(settings, a, b)
    assert len(res.outputs) == 2
    assert res.mapping.name == "a_ve_1_dosya_eslestirme.xlsx"
    txt = res.outputs[0].read_text(encoding="utf-8")
    ws = load_workbook(res.outputs[1])["Şikayetler"]
    code = ws["B2"].value  # Ahmet Yılmaz
    assert code in txt


def test_session_keeps_codes_for_next_file(settings, tmp_path):
    a = make_txt(tmp_path / "a.txt")
    s = Session(settings)
    r1 = s.analyze([a])
    s.save(r1)
    b = tmp_path / "b.txt"
    b.write_text("Yeni bir kişi: Kerem Doğan. Ahmet Yılmaz da vardı.", encoding="utf-8")
    r2 = s.analyze([b])
    res = s.save(r2)
    out = res.outputs[0].read_text(encoding="utf-8")
    assert "[KİŞİ_1] da vardı" in out  # Ahmet Yılmaz ilk dosyada KİŞİ_1 idi
    assert "Kerem" not in out and "[KİŞİ_3]" in out


def test_manual_add_and_type_change(settings, tmp_path):
    src = make_txt(tmp_path / "not.txt")
    s = Session(settings)
    result = s.analyze([src])
    c = s.add_manual(result, "Toplantı notu", T.OZEL)
    assert c.count == 1 and c.code == "[ÖZEL_1]"
    birim = next(x for x in result.candidates if x.etype == T.BIRIM)
    s.change_type(result, birim, T.KURUM)
    assert birim.etype == T.KURUM and birim.code == "[KURUM_1]"


def test_code_formats(settings, tmp_path):
    settings.code_format = "letter"
    src = make_txt(tmp_path / "not.txt")
    _, _, res = run(settings, src)
    out = res.outputs[0].read_text(encoding="utf-8")
    assert "Kişi A" in out and "Kişi B" in out
    back, _ = reverse_text(out, res.mapping)
    assert "Ahmet Yılmaz" in back


def test_chart_and_pivot_caches_are_cleaned(settings, tmp_path):
    src = make_xlsx(tmp_path / "grafik.xlsx")
    chart = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart><c:title><c:tx><c:rich>'
        "<a:p><a:r><a:t>Ahmet Yılmaz şikâyetleri</a:t></a:r></a:p></c:rich></c:tx></c:title>"
        "<c:cat><c:strRef><c:strCache><c:pt idx=\"0\"><c:v>Ayşe Kaya</c:v></c:pt></c:strCache></c:strRef></c:cat>"
        "</c:chart></c:chartSpace>"
    )
    pivot = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><cacheFields>'
        '<cacheField name="Ad Soyad"><sharedItems><s v="Ahmet Yılmaz"/><s v="AYŞE KAYA"/></sharedItems></cacheField>'
        "</cacheFields></pivotCacheDefinition>"
    )
    with zipfile.ZipFile(src, "a") as z:
        z.writestr("xl/charts/chart1.xml", chart)
        z.writestr("xl/pivotCache/pivotCacheDefinition1.xml", pivot)
    _, result, res = run(settings, src)
    c = zip_text(res.outputs[0], "xl/charts/chart1.xml")
    p = zip_text(res.outputs[0], "xl/pivotCache/pivotCacheDefinition1.xml")
    assert "Ahmet" not in c and "Ayşe" not in c and "[KİŞİ_1] şikâyetleri" in c
    assert "Ahmet" not in p and "AYŞE" not in p
